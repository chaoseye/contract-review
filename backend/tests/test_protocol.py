"""Генератор протокола: формат по documentation/PROTOCOL_FORMAT.md."""

import io
import json
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from PIL import Image, ImageDraw

from app.cli import main as cli_main
from app.documents.convert import find_soffice
from app.protocol import assets
from app.protocol.builder_docx import HEADER_FILL, TEXT_WIDTH, build_docx
from app.protocol.model import Letterhead, ProtocolData, ProtocolImages, ProtocolRow
from app.protocol.templates import (
    NBSP,
    ContractParty,
    ProtocolDraft,
    build_protocol_data,
    clause_label,
    contract_title,
    default_closing_items,
    long_date,
    named_as,
    party_intro,
)


def png(size: tuple[int, int], draw) -> bytes:
    image = Image.new("RGBA", size, (255, 255, 255, 0))
    draw(ImageDraw.Draw(image))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


SIGNATURE = png((400, 240), lambda d: d.line([(20, 200), (380, 40)], fill=(40, 40, 140, 255), width=6))
STAMP = png((400, 400), lambda d: d.ellipse([10, 10, 390, 390], outline=(40, 80, 200, 255), width=12))
HEADER = png((1100, 120), lambda d: d.rectangle([0, 0, 1099, 119], fill=(130, 130, 120, 255)))
# Бумага скана — не чисто белая.
PAPER = (247, 246, 242)


def scan(size: tuple[int, int], draw, paper=PAPER, orientation: int | None = None) -> bytes:
    """Скан без прозрачности: JPEG, как его сохраняет сканер или телефон."""
    image = Image.new("RGB", size, paper)
    draw(ImageDraw.Draw(image))
    exif = Image.Exif()
    if orientation:
        exif[0x0112] = orientation  # Orientation
    buffer = io.BytesIO()
    image.save(buffer, format="JPEG", quality=90, exif=exif)
    return buffer.getvalue()


SCAN_SIGNATURE = scan((400, 240), lambda d: d.line([(20, 200), (380, 40)], fill=(40, 40, 140), width=6))
SCAN_STAMP = scan((400, 400), lambda d: d.ellipse([10, 10, 390, 390], outline=(40, 80, 200), width=12))

ROWS = [
    ProtocolRow(
        clause=clause_label(["9.9"], new_clause=True),
        buyer_text="Пункт отсутствует.",
        supplier_text="9.9. Совокупный размер ответственности Поставщика не превышает 10% цены Договора.\n"
        "Упущенная выгода и косвенные убытки возмещению не подлежат.",
    ),
    ProtocolRow(
        clause=clause_label(["9.2"]),
        buyer_text="9.2. Неустойка 0,5% в день, но не более 20%.",
        supplier_text="9.2. Неустойка 0,1% в день, но не более 10%.",
        agreed_text="Редакция Поставщика",
    ),
    ProtocolRow(
        clause=clause_label(["6.1", "6.3", "6.5"]), buyer_text="6.1. …\n6.3. …", supplier_text="6.1. …"
    ),
]


def draft(**overrides) -> ProtocolDraft:
    values = {
        "contract_kind": "поставки",
        "contract_number": "12/2026",
        "contract_date": date(2026, 7, 31),
        "city": "г. Москва",
        "buyer": ContractParty(
            full_name="Акционерное общество «Пример»",
            short_name="АО «Пример»",
            representative="генерального директора Иванова И.И.",
            basis="Устава",
        ),
        "supplier": ContractParty(
            full_name="Индивидуальный предприниматель Петров Пётр Петрович",
            short_name="ИП Петров Пётр Петрович",
            signer="П.П. Петров",
        ),
        "rows": ROWS,
    }
    values.update(overrides)
    return ProtocolDraft(**values)


def open_docx(data: bytes):
    return Document(io.BytesIO(data))


def anchors(document) -> list:
    return document.element.body.findall(f".//{qn('wp:anchor')}")


# --- шаблоны ------------------------------------------------------------------------------


def test_texts_and_labels() -> None:
    assert long_date(date(2026, 7, 31)) == f"31{NBSP}июля 2026{NBSP}г."
    assert clause_label(["9.2"]) == f"п.{NBSP}9.2"
    assert clause_label(["6.1", "6.3", "6.5"]) == f"пп.{NBSP}6.1, 6.3 и 6.5"
    assert clause_label(["9.9."], new_clause=True) == f"п.{NBSP}9.9 (новый пункт)"
    assert clause_label([]) == ""
    assert (
        contract_title("поставки", "12/2026", date(2026, 7, 31))
        == f"Договору поставки №{NBSP}12/2026 от 31.07.2026"
    )
    assert contract_title("поставки", "", None) == f"Договору поставки №{NBSP}____"


@pytest.mark.parametrize(
    ("name", "expected"),
    [
        ("Акционерное общество «Пример»", "именуемое"),
        ("ООО «Ромашка»", "именуемое"),
        ("Индивидуальный предприниматель Петров Пётр Петрович", "именуемый"),
        ("Индивидуальный предприниматель Иванова Анна Сергеевна", "именуемая"),
        ("Компания «Альфа»", "именуемая"),
    ],
)
def test_named_as(name: str, expected: str) -> None:
    assert named_as(name) == expected


def test_party_intro_and_closing_items() -> None:
    assert party_intro("АО «Пример»", "Покупатель", "директора Иванова И.И.", "Устава") == (
        "АО «Пример», именуемое в дальнейшем «Покупатель», в лице директора Иванова И.И., "
        "действующего на основании Устава"
    )
    items = default_closing_items("поставки", "12/2026", ["Спецификация № 1 на 2 листах"])
    assert len(items) == 4
    assert items[0].startswith(f"Остальные условия Договора поставки №{NBSP}12/2026")
    assert items[3] == "Приложения: Спецификация № 1 на 2 листах."


def test_build_protocol_data_respects_party_order() -> None:
    buyer_first = build_protocol_data(draft())
    assert buyer_first.preamble.startswith(
        "Акционерное общество «Пример», именуемое в дальнейшем «Покупатель»"
    )
    assert "составили настоящий Протокол разногласий к Договору поставки" in buyer_first.preamble
    assert buyer_first.subtitle == f"к Договору поставки №{NBSP}12/2026 от 31.07.2026"
    assert buyer_first.date == date(2026, 7, 31)
    assert buyer_first.supplier.name == "ИП Петров Пётр Петрович"
    supplier_first = build_protocol_data(draft(supplier_first=True, protocol_date=date(2026, 8, 3)))
    assert supplier_first.preamble.startswith("Индивидуальный предприниматель Петров")
    assert supplier_first.date == date(2026, 8, 3)


# --- документ -----------------------------------------------------------------------------


def test_page_fonts_and_title_block() -> None:
    document = open_docx(build_docx(build_protocol_data(draft())))
    section = document.sections[0]
    assert section.orientation == WD_ORIENT.LANDSCAPE
    assert abs(section.page_width - Mm(297)) < Mm(1) and abs(section.page_height - Mm(210)) < Mm(1)
    assert abs(section.left_margin - Mm(20)) < Mm(0.5)

    normal = document.styles["Normal"]
    fonts = normal.element.rPr.find(qn("w:rFonts"))
    assert {fonts.get(qn(a)) for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia")} == {"Times New Roman"}
    assert not [a for a in fonts.attrib if a.endswith("Theme")]
    assert normal.element.rPr.find(qn("w:lang")).get(qn("w:val")) == "ru-RU"

    texts = [p.text for p in document.paragraphs if p.text]
    assert texts[0] == "ПРОТОКОЛ РАЗНОГЛАСИЙ"
    assert texts[1] == f"к Договору поставки №{NBSP}12/2026 от 31.07.2026"
    assert texts[2] == f"г. Москва\t31{NBSP}июля 2026{NBSP}г."
    assert texts[3].startswith("Акционерное общество «Пример»")
    assert document.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert document.paragraphs[0].runs[0].bold
    assert document.core_properties.author == "ИП Петров Пётр Петрович"


def test_rows_table() -> None:
    document = open_docx(build_docx(build_protocol_data(draft())))
    table = document.tables[0]
    assert len(table.columns) == 5
    assert [c.text for c in table.rows[0].cells] == [
        "№",
        "Пункт Договора",
        "Редакция Покупателя",
        "Редакция Поставщика",
        "Согласованная редакция",
    ]
    header_tr = table.rows[0]._tr
    assert header_tr.trPr.find(qn("w:tblHeader")) is not None
    assert header_tr.trPr.find(qn("w:cantSplit")) is not None
    fill = table.rows[0].cells[0]._tc.tcPr.find(qn("w:shd")).get(qn("w:fill"))
    assert fill == HEADER_FILL
    assert table._tbl.tblPr.find(qn("w:tblLayout")).get(qn("w:type")) == "fixed"

    grid = [int(col.get(qn("w:w"))) for col in table._tbl.tblGrid.findall(qn("w:gridCol"))]
    total = TEXT_WIDTH / 635
    assert [round(width / total * 100, 1) for width in grid] == [3.5, 7.5, 23.5, 36.0, 29.5]

    first, second, third = (table.rows[i] for i in (1, 2, 3))
    assert [c.text for c in first.cells[:2]] == ["1", f"п.{NBSP}9.9 (новый пункт)"]
    assert [p.text for p in first.cells[3].paragraphs] == [
        "9.9. Совокупный размер ответственности Поставщика не превышает 10% цены Договора.",
        "Упущенная выгода и косвенные убытки возмещению не подлежат.",
    ]
    agreed = second.cells[4].paragraphs[0]
    assert agreed.text == "Редакция Поставщика"
    assert agreed.alignment == WD_ALIGN_PARAGRAPH.CENTER and agreed.runs[0].bold
    assert third.cells[1].text == f"пп.{NBSP}6.1, 6.3 и 6.5"
    assert third.cells[4].text == ""
    assert first.cells[3].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.JUSTIFY


def test_table_font_size_is_configurable() -> None:
    data = build_protocol_data(draft())
    data.table_font_size = 9
    table = open_docx(build_docx(data)).tables[0]
    sizes = {
        run.font.size.pt
        for row in table.rows
        for cell in row.cells
        for p in cell.paragraphs
        for run in p.runs
    }
    assert sizes == {9.0}


def test_table_without_agreed_column() -> None:
    document = open_docx(build_docx(build_protocol_data(draft(include_agreed_column=False))))
    table = document.tables[0]
    assert len(table.columns) == 4
    assert table.rows[0].cells[-1].text == "Редакция Поставщика"


TBL_PR_ORDER = ["w:tblStyle", "w:tblW", "w:jc", "w:tblLayout", "w:tblCellMar", "w:tblLook"]


def test_table_properties_follow_schema_order() -> None:
    """Word считает документ повреждённым, если порядок элементов w:tblPr нарушен."""
    document = open_docx(build_docx(build_protocol_data(draft())))
    for table in document.tables:
        tags = [child.tag for child in table._tbl.tblPr]
        positions = [tags.index(qn(tag)) for tag in TBL_PR_ORDER if qn(tag) in tags]
        assert positions == sorted(positions)
        assert len(tags) == len(set(tags))


def test_closing_items_and_signature_block() -> None:
    document = open_docx(build_docx(build_protocol_data(draft())))
    paragraphs = [p for p in document.paragraphs if p.text]
    closing_index = next(i for i, p in enumerate(paragraphs) if p.text == "Заключительные положения")
    items = paragraphs[closing_index + 1 : closing_index + 4]
    assert [p.text[:3] for p in items] == ["1. ", "2. ", "3. "]
    assert items[-1].paragraph_format.keep_with_next
    assert not items[0].paragraph_format.keep_with_next

    buyer, supplier = document.tables[1].rows[0].cells
    assert [p.text for p in buyer.paragraphs if p.text] == [
        "ПОКУПАТЕЛЬ:",
        "АО «Пример»",
        "________________ / ________________ /",
        "М.П.",
    ]
    assert [p.text for p in supplier.paragraphs if p.text] == [
        "ПОСТАВЩИК:",
        "ИП Петров Пётр Петрович",
        "________________ / П.П. Петров /",
        "М.П.",
    ]
    assert document.tables[1].rows[0]._tr.trPr.find(qn("w:cantSplit")) is not None


def test_footer_page_number_and_no_letterhead_by_default() -> None:
    document = open_docx(build_docx(build_protocol_data(draft()), ProtocolImages(header=HEADER)))
    section = document.sections[0]
    assert not section.different_first_page_header_footer
    page = section.footer.paragraphs[0]
    assert [node.text.strip() for node in page._p.iter(qn("w:instrText"))] == ["PAGE"]
    kinds = [node.get(qn("w:fldCharType")) for node in page._p.iter(qn("w:fldChar"))]
    assert kinds == ["begin", "separate", "end"]
    # Номер — 9 pt, как слово «Страница»: у каждой части поля свой размер, иначе Word наберёт 11 pt.
    assert page.text == "Страница 1"
    assert {run.font.size for run in page.runs} == {Pt(9)}


def test_letterhead_on_first_page_and_contacts_in_footer() -> None:
    data = build_protocol_data(draft(letterhead=Letterhead(enabled=True, footer="г. Пример, тел. 000")))
    document = open_docx(build_docx(data, ProtocolImages(header=HEADER)))
    section = document.sections[0]
    assert section.different_first_page_header_footer
    assert section.first_page_header._element.findall(f".//{qn('a:blip')}")
    for footer in (section.first_page_footer, section.footer):
        # Номер страницы — поле PAGE; в .text попадает его сохранённое значение «1».
        assert [p.text for p in footer.paragraphs] == ["г. Пример, тел. 000", "Страница 1"]
        assert footer.paragraphs[1]._p.find(f".//{qn('w:instrText')}") is not None


def test_no_facsimile_without_flags() -> None:
    images = ProtocolImages(signature=SIGNATURE, stamp=STAMP)
    assert anchors(open_docx(build_docx(build_protocol_data(draft()), images))) == []


@pytest.mark.parametrize(
    ("with_signature", "with_stamp", "count"), [(True, False, 1), (False, True, 1), (True, True, 2)]
)
def test_facsimile_is_floating_in_supplier_block(with_signature: bool, with_stamp: bool, count: int) -> None:
    images = ProtocolImages(signature=SIGNATURE, stamp=STAMP)
    document = open_docx(
        build_docx(build_protocol_data(draft()), images, with_signature=with_signature, with_stamp=with_stamp)
    )
    found = anchors(document)
    assert len(found) == count
    supplier_cell = document.tables[1].rows[0].cells[1]._tc
    for anchor in found:
        assert anchor.getparent() is not None
        assert supplier_cell in list(anchor.iterancestors())
        assert anchor.find(qn("wp:wrapNone")) is not None
        assert anchor.get("behindDoc") == "0"
        assert anchor.find(f"{qn('wp:positionH')}/{qn('wp:posOffset')}") is not None
    assert not document.element.body.findall(f".//{qn('wp:inline')}")
    if count == 2:
        heights = [int(anchor.get("relativeHeight")) for anchor in found]
        assert heights[1] > heights[0]  # печать поверх подписи


def test_stamp_is_rotated_and_semi_transparent() -> None:
    prepared, ratio = assets.prepare_stamp(STAMP, "12/2026")
    image = Image.open(io.BytesIO(prepared))
    assert image.mode == "RGBA"
    assert image.getchannel("A").getextrema()[1] <= int(255 * assets.STAMP_OPACITY) + 1
    assert 0.9 < ratio < 1.1
    again, _ = assets.prepare_stamp(STAMP, "12/2026")
    assert again == prepared  # повторный выпуск — тот же оттиск
    angles = {round(assets.stamp_angle(seed), 3) for seed in ("1", "2", "3", "12/2026")}
    assert all(assets.STAMP_MIN_ANGLE <= angle <= assets.STAMP_MAX_ANGLE for angle in angles)
    assert len(angles) > 1


def test_signature_is_trimmed() -> None:
    prepared, ratio = assets.prepare_signature(SIGNATURE)
    image = Image.open(io.BytesIO(prepared))
    assert image.size[0] < 400 and image.size[1] < 240
    assert abs(ratio - image.height / image.width) < 1e-6


def test_scan_paper_becomes_transparent() -> None:
    """Бумага скана прозрачна, чернила — по своей густоте; поверх белого — как скан."""
    inks = [(40, 40, 140), (20, 20, 20), (150, 170, 220), (230, 232, 236)]
    source = Image.new("RGB", (len(inks) + 1, 1), PAPER)
    for x, color in enumerate(inks):
        source.putpixel((x, 0), color)
    buffer = io.BytesIO()
    source.save(buffer, format="PNG")  # PNG без прозрачности обрабатывается так же, как JPEG
    layer = assets._paper_to_alpha(assets._load_rgba(buffer.getvalue()))
    assert layer.getpixel((len(inks), 0))[3] == 0  # бумага
    assert layer.getpixel((3, 0))[3] <= assets.FAINT_ALPHA  # почти белое — тоже фон
    assert layer.getpixel((1, 0))[3] > 200  # чёрные чернила почти непрозрачны
    over_white = Image.alpha_composite(Image.new("RGBA", layer.size, "white"), layer).convert("RGB")
    for x, color in enumerate(inks):
        expected = [min(255, value * 255 // assets.PAPER_LEVEL) for value in color]
        assert all(abs(a - b) <= 2 for a, b in zip(over_white.getpixel((x, 0)), expected, strict=True))


def test_scan_signature_and_stamp_get_transparent_background() -> None:
    signature, ratio = assets.prepare_signature(SCAN_SIGNATURE)
    image = Image.open(io.BytesIO(signature))
    # Бумага стала прозрачной, поэтому поля обрезаны, как у PNG с прозрачным фоном.
    assert image.mode == "RGBA" and image.size[0] < 400 and image.size[1] < 240
    assert image.getpixel((0, 0))[3] <= assets.FAINT_ALPHA  # угол напротив росчерка — бумага
    assert image.getchannel("A").getextrema()[1] > 200
    assert abs(ratio - image.height / image.width) < 1e-6

    stamp, _ = assets.prepare_stamp(SCAN_STAMP, "12/2026")
    image = Image.open(io.BytesIO(stamp))
    assert image.getpixel((image.width // 2, image.height // 2))[3] == 0  # внутри оттиска — бумага
    assert 0 < image.getchannel("A").getextrema()[1] <= int(255 * assets.STAMP_OPACITY) + 1


def test_image_with_transparency_is_kept() -> None:
    # Прозрачность задана — картинку не трогаем, даже белые места в ней.
    framed = png((200, 100), lambda d: d.rectangle([20, 20, 180, 80], fill=(255, 255, 255, 255)))
    prepared, _ = assets.prepare_signature(framed)
    image = Image.open(io.BytesIO(prepared))
    assert image.size == (161, 61) and image.getchannel("A").getextrema() == (255, 255)
    assert image.getpixel((80, 30)) == (255, 255, 255, 255)


def test_background_that_covers_text() -> None:
    def stroke(draw) -> None:
        draw.line([(20, 200), (380, 40)], fill=(40, 40, 140), width=6)

    def opaque_png(color: tuple[int, int, int]) -> bytes:
        return png((200, 100), lambda d: d.rectangle([0, 0, 199, 99], fill=(*color, 255)))

    assert not assets.covers_text(SIGNATURE)
    assert not assets.covers_text(SCAN_SIGNATURE)
    assert not assets.covers_text(SCAN_STAMP)
    # Фото на серой бумаге: серое останется и закроет линию подписи.
    assert assets.covers_text(scan((400, 240), stroke, paper=(190, 185, 175)))
    # PNG с непрозрачным цветным фоном; белый непрозрачный фон станет прозрачным, как у скана.
    assert assets.covers_text(opaque_png((200, 220, 255)))
    assert not assets.covers_text(opaque_png((255, 255, 255)))
    # Росчерк, который касается края, фоном не считается.
    touching = scan((400, 240), lambda d: d.line([(0, 239), (399, 0)], fill=(40, 40, 140), width=8))
    assert not assets.covers_text(touching)


def test_exif_orientation_is_applied() -> None:
    # Снимок телефона: пиксели лёжа, в EXIF — «повернуть на 90°». Чернила слева, после поворота — сверху.
    photo = scan((400, 200), lambda d: d.rectangle([0, 0, 99, 199], fill=(20, 20, 20)), orientation=6)
    prepared, ratio = assets.prepare_signature(photo)
    width, height = Image.open(io.BytesIO(prepared)).size
    # Без поворота было бы около 100 × 200; лишняя строка-другая — ореол JPEG у края чернил.
    assert width == 200 and 100 <= height <= 104
    assert abs(ratio - height / width) < 1e-6


def test_large_scan_is_read_reduced() -> None:
    # Большой JPEG читается сразу уменьшенным; сторона в протоколе — не больше MAX_SIDE_PX.
    large = scan((2600, 1200), lambda d: d.line([(100, 1100), (2500, 100)], fill=(40, 40, 140), width=24))
    prepared, ratio = assets.prepare_signature(large)
    image = Image.open(io.BytesIO(prepared))
    assert max(image.size) <= assets.MAX_SIDE_PX and image.getpixel((0, 0))[3] <= assets.FAINT_ALPHA
    assert abs(ratio - 1000 / 2400) < 0.03
    assert not assets.covers_text(large)


def test_scan_facsimile_goes_into_docx_transparent() -> None:
    images = ProtocolImages(signature=SCAN_SIGNATURE, stamp=SCAN_STAMP)
    data = build_protocol_data(draft())
    document = open_docx(build_docx(data, images, with_signature=True, with_stamp=True))
    assert len(anchors(document)) == 2
    pictures = [part for part in document.part.package.iter_parts() if part.content_type == "image/png"]
    assert len(pictures) == 2
    for part in pictures:
        with Image.open(io.BytesIO(part.blob)) as picture:
            assert picture.mode == "RGBA" and picture.getchannel("A").getextrema()[0] == 0


def test_cli_builds_protocol_from_json(tmp_path: Path) -> None:
    data = build_protocol_data(draft())
    source = tmp_path / "data.json"
    source.write_text(data.model_dump_json(), encoding="utf-8")
    signature = tmp_path / "signature.png"
    signature.write_bytes(SIGNATURE)
    out = tmp_path / "protocol.docx"
    cli_main(["protocol", str(source), "--out", str(out), "--signature", str(signature)])
    document = open_docx(out.read_bytes())
    assert document.paragraphs[0].text == "ПРОТОКОЛ РАЗНОГЛАСИЙ"
    assert len(anchors(document)) == 1


def test_json_round_trip() -> None:
    data = build_protocol_data(draft())
    restored = ProtocolData.model_validate(json.loads(data.model_dump_json()))
    assert restored == data


@pytest.mark.skipif(find_soffice() is None, reason="нет LibreOffice (проверяется в Docker)")
def test_pdf_is_built_from_docx() -> None:
    import pdfplumber

    from app.protocol.pdf import docx_to_pdf

    data = build_protocol_data(draft(letterhead=Letterhead(enabled=True, footer="г. Пример, тел. 000")))
    pdf = docx_to_pdf(
        build_docx(
            data,
            ProtocolImages(header=HEADER, signature=SIGNATURE, stamp=STAMP),
            with_signature=True,
            with_stamp=True,
        )
    )
    with pdfplumber.open(io.BytesIO(pdf)) as document:
        text = "\n".join(page.extract_text() or "" for page in document.pages)
        assert abs(document.pages[0].width - 842) < 2 and abs(document.pages[0].height - 595) < 2
        images = sum(len(page.images) for page in document.pages)
    assert "ПРОТОКОЛ РАЗНОГЛАСИЙ" in text
    assert "Редакция Поставщика" in text
    assert "Страница 1" in text
    assert images >= 3  # шапка, подпись, печать
