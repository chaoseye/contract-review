"""Этап 4, серверная часть: протокол в проверке, скачивание, организация, журнал, ограничение входа."""

import copy
import io
import re
from datetime import UTC, date, datetime
from urllib.parse import unquote

import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image
from sqlalchemy import select

from app import cli
from app.analysis import profiles
from app.api import protocol as protocol_api
from app.api.ratelimit import LoginLimiter
from app.documents.convert import find_soffice
from app.models import AuditEntry, Check, Organization, UserRole
from app.protocol.editing import (
    ProtocolState,
    draft_from_result,
    file_name,
    preamble_preview,
    to_protocol_data,
)
from tests.analysis_helpers import BLOCKS, FakeProvider, answer, risk
from tests.conftest import PASSWORD, add_user
from tests.test_analysis_job import add_profile, contract_answer, uploaded
from tests.test_checks_api import login, run_next_job

WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"


def png() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGBA", (60, 30), (20, 40, 160, 255)).save(buffer, format="PNG")
    return buffer.getvalue()


def add_organization(db_factory, signature: bool = True, stamp: bool = True) -> None:
    with db_factory() as db:
        db.add(
            Organization(
                full_name="Индивидуальный предприниматель Петров Пётр Петрович",
                short_name="ИП Петров Пётр Петрович",
                signer="П.П. Петров",
                representative="",
                basis="",
                default_city="г. Пример",
                requisites={"inn": "770000000156"},
                footer_text="",
                signature_png=png() if signature else None,
                stamp_png=png() if stamp else None,
            )
        )
        db.commit()


@pytest.fixture
def provider(monkeypatch) -> FakeProvider:
    fake = FakeProvider()
    monkeypatch.setattr(profiles, "build_provider", lambda profile, settings, **kwargs: fake)
    return fake


@pytest.fixture
def ready_check(client, db_factory, settings, tmp_path, provider) -> str:
    """Проверка в статусе ready: менеджер вошёл, договор прочитан и проанализирован."""
    add_user(db_factory)
    add_profile(db_factory)
    login(client, "manager")
    provider.responses.append(contract_answer())
    check_id = uploaded(client, db_factory, settings, tmp_path)
    run_next_job(db_factory, settings)
    assert client.get(f"/api/checks/{check_id}").json()["status"] == "ready"
    return check_id


def audit_rows(db_factory) -> list[AuditEntry]:
    with db_factory() as db:
        return list(db.scalars(select(AuditEntry).order_by(AuditEntry.id)))


# --- черновик протокола ------------------------------------------------------------------------


def test_draft_from_analysis() -> None:
    result = {
        "answer": answer(
            risks=[
                risk("R1", clauses=["9.8"], block_ids=["B0003"], quote="удержать сумму неустойки"),
                risk("R2", clauses=["9.2"]),
                risk("R3", clauses=["9.9"], is_new_clause=True, quote="", block_ids=[]),
                risk("R4", clauses=["15.1"], quote="текст пункта 15.1", severity="low"),
                risk("R5", clauses=["2.4"], proposed_wording=None),
            ],
            contract={**answer()["contract"], "city": None, "parties_order": ["supplier", "buyer"]},
        )
    }
    state = draft_from_result(result, BLOCKS, None, date(2026, 10, 16))
    assert [row.risk_id for row in state.rows] == ["R5", "R2", "R1", "R3", "R4"]
    by_id = {row.risk_id: row for row in state.rows}
    # Колонка покупателя — полный текст пункта из договора, а не цитата модели.
    assert by_id["R1"].buyer_text == BLOCKS[2]["text"]
    assert by_id["R1"].clause == "п.\u00a09.8"
    assert by_id["R3"].buyer_text == "Пункт отсутствует."
    assert by_id["R3"].clause == "п.\u00a09.9 (новый пункт)"
    # Пункта 15.1 в разметке нет — берётся цитата.
    assert by_id["R4"].buyer_text == "текст пункта 15.1"
    assert [row.include for row in state.rows] == [False, True, True, True, False]
    meta = state.meta
    assert (meta.contract_number, meta.contract_date, meta.protocol_date) == (
        "12/2026",
        date(2026, 7, 31),
        date(2026, 7, 31),
    )
    assert meta.buyer_full_name == "АО «Пример»" and meta.supplier_first is True and meta.city == ""


def test_draft_fits_long_values_from_model() -> None:
    # Раньше длинное поле ответа давало ValidationError и вкладка «Протокол» отвечала 500.
    long_wording = "9.2. " + "Поставщик уплачивает неустойку. " * 700
    result = {
        "answer": answer(
            risks=[
                risk("R" * 21, proposed_wording=long_wording),
                risk("R2", clauses=[f"{index}.1" for index in range(1, 80)], quote="ц" * 30_000),
                # id — ссылка на риск: длинные и похожие не обрезаются и не сливаются.
                risk("R" + "1" * 24 + "A"),
                risk("R" + "1" * 24 + "B"),
                risk("Z" * 201),
                *(risk(f"X{index}", clauses=["9.2"]) for index in range(320)),
            ],
            contract={**answer()["contract"], "number": "№" * 150, "buyer": {"name": "Б" * 600}},
        )
    }
    state = draft_from_result(result, BLOCKS, None, date(2026, 10, 16))
    assert len(state.meta.contract_number) <= 100 and state.meta.contract_number.endswith("…")
    assert len(state.meta.buyer_full_name) <= 500
    assert len(state.rows) == 300
    by_id = {row.risk_id: row for row in state.rows}
    long_row = by_id["R" * 21]
    assert len(long_row.supplier_text) <= 20000 and long_row.supplier_text.endswith("…")
    # Обрезанную редакцию менеджер сначала проверяет: в протокол по умолчанию она не идёт.
    assert long_row.include is False
    many = by_id["R2"]
    assert len(many.clause) <= 200 and many.clause.endswith("…")
    assert len(many.buyer_text) <= 20000 and many.include is True
    assert {"R" + "1" * 24 + "A", "R" + "1" * 24 + "B"} <= set(by_id)
    # Совсем длинный id — строка без ссылки на риск, как свой пункт менеджера.
    assert None in by_id and not any(key and key.startswith("Z") for key in by_id)
    # Черновик проходит ту же проверку, что и сохранённые правки.
    assert type(state).model_validate(state.model_dump()) == state


def test_protocol_data_and_file_name() -> None:
    state = draft_from_result({"answer": answer()}, BLOCKS, None, date(2026, 10, 16))
    organization = Organization(
        full_name="Индивидуальный предприниматель Петров Пётр Петрович",
        short_name="ИП Петров П.П.",
        signer="П.П. Петров",
        representative="",
        basis="",
    )
    with pytest.raises(ValueError, match="данные нашей организации"):
        to_protocol_data(state, None)
    data = to_protocol_data(state, organization)
    assert data.supplier.name == "ИП Петров П.П." and data.supplier.signer == "П.П. Петров"
    assert "именуемый в дальнейшем «Поставщик»" in data.preamble
    assert [row.clause for row in data.rows] == ["п.\u00a09.2"]
    state.meta.preamble = "Своё вступление."
    state.meta.table_font_size = 9
    data = to_protocol_data(state, organization)
    assert data.preamble == "Своё вступление." and data.table_font_size == 9
    assert preamble_preview(state, organization) == "Своё вступление."
    state.meta.contract_number = 'А/12:"2026"'
    assert file_name(state, "docx") == "Протокол разногласий к договору № А-12-2026.docx"
    state.meta.contract_number = ""
    assert file_name(state, "pdf") == "Протокол разногласий.pdf"
    state.rows[0].include = False
    with pytest.raises(ValueError, match="нет ни одной строки"):
        to_protocol_data(state, organization)
    state.rows[0].include = True
    state.meta.buyer_full_name = " "
    with pytest.raises(ValueError, match="наименование покупателя"):
        to_protocol_data(state, organization)


# --- API протокола -----------------------------------------------------------------------------


def test_protocol_requires_analysis(client, db_factory, settings, tmp_path) -> None:
    add_user(db_factory)
    login(client, "manager")
    check_id = uploaded(client, db_factory, settings, tmp_path)
    response = client.get(f"/api/checks/{check_id}/protocol")
    assert response.status_code == 409


def test_protocol_draft_edit_and_download(client, db_factory, ready_check) -> None:
    add_organization(db_factory)
    body = client.get(f"/api/checks/{ready_check}/protocol").json()
    assert body["saved"] is False
    assert body["organization"] == {"ready": True, "signature": True, "stamp": True}
    state = body["state"]
    assert state["meta"]["buyer_full_name"] == "АО «Пример»"
    assert state["meta"]["city"] == "г. Пример"  # в договоре города нет — город организации
    [row] = state["rows"]
    assert row["risk_id"] == "R1" and row["include"] is True
    assert row["buyer_text"].startswith("9.2. В случае просрочки поставки")
    assert "Индивидуальный предприниматель Петров Пётр Петрович" in body["preamble_preview"]

    state["rows"].append(
        {
            "include": True,
            "clause": "п. 14.5",
            "buyer_text": "Пункт отсутствует.",
            "supplier_text": "14.5. Споры — в суде по месту ответчика.",
        }
    )
    state["rows"].reverse()
    state["meta"]["buyer_signer"] = "И.И. Иванов"
    saved = client.put(f"/api/checks/{ready_check}/protocol", json=state).json()
    assert saved["saved"] is True
    assert [row["clause"] for row in saved["state"]["rows"]] == ["п. 14.5", "п.\u00a09.2"]
    assert client.get(f"/api/checks/{ready_check}/protocol").json()["state"] == saved["state"]

    response = client.get(f"/api/checks/{ready_check}/protocol.docx")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    disposition = response.headers["content-disposition"]
    assert (
        unquote(disposition.split("filename*=UTF-8''")[1]) == "Протокол разногласий к договору № 12-2026.docx"
    )
    document = Document(io.BytesIO(response.content))
    table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "14.5. Споры — в суде по месту ответчика." in table_text
    assert "ИП Петров Пётр Петрович" in table_text and "И.И. Иванов" in table_text
    assert document.element.body.findall(f".//{WP}anchor") == []

    signed = client.get(f"/api/checks/{ready_check}/protocol.docx?signature=true&stamp=true")
    assert signed.status_code == 200
    assert len(Document(io.BytesIO(signed.content)).element.body.findall(f".//{WP}anchor")) == 2

    downloads = [entry for entry in audit_rows(db_factory) if entry.action == "protocol_docx"]
    assert [(entry.with_signature, entry.with_stamp) for entry in downloads] == [(False, False), (True, True)]
    assert all(
        entry.counterparty == "АО «Пример»" and entry.contract_number == "12/2026" for entry in downloads
    )


def test_download_refuses_missing_data(client, db_factory, ready_check) -> None:
    response = client.get(f"/api/checks/{ready_check}/protocol.docx")
    assert response.status_code == 409 and "данные нашей организации" in response.json()["detail"]
    assert "полное наименование" in response.json()["detail"]
    add_organization(db_factory, signature=False)
    response = client.get(f"/api/checks/{ready_check}/protocol.docx?signature=true")
    assert response.status_code == 409 and "нет подписи" in response.json()["detail"]
    state = client.get(f"/api/checks/{ready_check}/protocol").json()["state"]
    state["rows"][0]["include"] = False
    client.put(f"/api/checks/{ready_check}/protocol", json=state)
    response = client.get(f"/api/checks/{ready_check}/protocol.docx")
    assert response.status_code == 409 and "нет ни одной строки" in response.json()["detail"]


def test_organization_needs_only_full_name() -> None:
    assert Organization(full_name="ИП Петров Пётр Петрович", short_name="", signer="").ready_for_protocol
    assert not Organization(full_name=" ", short_name="ИП Петров", signer="П.П. Петров").ready_for_protocol


def test_protocol_without_signer(client, db_factory, ready_check) -> None:
    # Подписант не обязателен: в блоке подписей поставщика линия, в подписях — полное наименование.
    with db_factory() as db:
        db.add(
            Organization(
                full_name="Индивидуальный предприниматель Петров Пётр Петрович",
                short_name="",
                signer="",
                representative="",
                basis="",
                default_city="",
                requisites={},
                footer_text="",
            )
        )
        db.commit()
    assert client.get(f"/api/checks/{ready_check}/protocol").json()["organization"]["ready"] is True
    response = client.get(f"/api/checks/{ready_check}/protocol.docx")
    assert response.status_code == 200
    supplier = Document(io.BytesIO(response.content)).tables[1].rows[0].cells[1]
    assert [p.text for p in supplier.paragraphs if p.text] == [
        "ПОСТАВЩИК:",
        "Индивидуальный предприниматель Петров Пётр Петрович",
        "________________ / ________________ /",
        "М.П.",
    ]


def test_protocol_date_is_local_today(client, db_factory, ready_check, monkeypatch) -> None:
    # Договор без даты: дата протокола — сегодня по часам браузера менеджера, а не по UTC.
    add_organization(db_factory)
    with db_factory() as db:
        check = db.get(Check, ready_check)
        result = copy.deepcopy(check.result)
        result["answer"]["contract"]["date"] = None
        check.result = result
        db.commit()
    # 00:30 по Москве (UTC+3) — в UTC ещё вчера.
    monkeypatch.setattr(protocol_api, "utcnow", lambda: datetime(2026, 10, 20, 21, 30, tzinfo=UTC))
    url = f"/api/checks/{ready_check}/protocol"

    def protocol_date(response) -> str:
        return response.json()["state"]["meta"]["protocol_date"]

    assert protocol_date(client.get(url)) == "2026-10-20"  # без смещения — UTC, как раньше
    assert protocol_date(client.get(url, params={"utc_offset": 180})) == "2026-10-21"
    assert protocol_date(client.get(url, params={"utc_offset": -300})) == "2026-10-20"
    assert protocol_date(client.post(f"{url}/reset", params={"utc_offset": 180})) == "2026-10-21"
    # Несохранённый черновик скачивается с той же датой.
    docx = client.get(f"{url}.docx", params={"utc_offset": 180})
    place = [p.text for p in Document(io.BytesIO(docx.content)).paragraphs if "\t" in p.text]
    assert place == ["г. Пример\t21\u00a0октября 2026\u00a0г."]
    # Сохранённая дата смещением не меняется.
    state = client.get(url, params={"utc_offset": 180}).json()["state"]
    client.put(url, json=state)
    assert protocol_date(client.get(url)) == "2026-10-21"
    assert client.get(url, params={"utc_offset": 15 * 60}).status_code == 422


def test_pdf_download(client, db_factory, ready_check) -> None:
    add_organization(db_factory)
    response = client.get(f"/api/checks/{ready_check}/protocol.pdf")
    if find_soffice() is None:
        assert response.status_code == 503
        assert response.json()["detail"].startswith("PDF не собран")
        assert not any(entry.action == "protocol_pdf" for entry in audit_rows(db_factory))
    else:  # pragma: no cover — в Docker с LibreOffice
        assert response.status_code == 200 and response.content.startswith(b"%PDF")


def test_reset_and_reanalysis_drop_edits(client, db_factory, ready_check, provider) -> None:
    add_organization(db_factory)
    state = client.get(f"/api/checks/{ready_check}/protocol").json()["state"]
    state["rows"][0]["supplier_text"] = "Моя редакция"
    client.put(f"/api/checks/{ready_check}/protocol", json=state)
    reset = client.post(f"/api/checks/{ready_check}/protocol/reset").json()
    assert reset["saved"] is False and reset["state"]["rows"][0]["supplier_text"] != "Моя редакция"

    client.put(f"/api/checks/{ready_check}/protocol", json=state)
    assert client.post(f"/api/checks/{ready_check}/analyze").status_code == 202
    with db_factory() as db:
        assert db.get(Check, ready_check).protocol is None


def test_protocol_permissions(app, client, db_factory, ready_check) -> None:
    add_organization(db_factory)
    add_user(db_factory, login="admin", role=UserRole.ADMIN)
    add_user(db_factory, login="other")
    state = client.get(f"/api/checks/{ready_check}/protocol").json()["state"]
    with TestClient(app) as admin, TestClient(app) as other:
        for http, name in ((admin, "admin"), (other, "other")):
            http.headers["X-Requested-With"] = "contract-review"
            login(http, name)
        assert other.get(f"/api/checks/{ready_check}/protocol").status_code == 404
        assert other.get(f"/api/checks/{ready_check}/protocol.docx").status_code == 404
        assert admin.get(f"/api/checks/{ready_check}/protocol").status_code == 200
        assert admin.put(f"/api/checks/{ready_check}/protocol", json=state).status_code == 403
        assert admin.get(f"/api/checks/{ready_check}/protocol.docx").status_code == 403
    bad = {**state, "rows": [{"clause": "x" * 300}]}
    assert client.put(f"/api/checks/{ready_check}/protocol", json=bad).status_code == 422


def test_list_shows_contract_risks_and_last_download(
    client, db_factory, settings, ready_check, provider
) -> None:
    add_organization(db_factory)
    [item] = client.get("/api/checks").json()
    contract = (item["contract_kind"], item["contract_number"], item["buyer_name"])
    assert contract == ("поставки", "12/2026", "АО «Пример»")
    assert (item["risk_count"], item["stop_count"]) == (1, 1)
    assert item["protocol_downloaded_at"] is None and item["protocol_format"] is None

    # Правки протокола важнее ответа модели — как в журнале.
    state = client.get(f"/api/checks/{ready_check}/protocol").json()["state"]
    state["meta"].update(contract_number="15/2026", buyer_full_name="ООО «Северный склад»")
    client.put(f"/api/checks/{ready_check}/protocol", json=state)
    assert client.get(f"/api/checks/{ready_check}/protocol.pdf").status_code in (200, 503)
    assert client.get(f"/api/checks/{ready_check}/protocol.docx").status_code == 200
    [item] = client.get("/api/checks").json()
    assert (item["contract_number"], item["buyer_name"]) == ("15/2026", "ООО «Северный склад»")
    assert item["protocol_format"] == "docx" and item["protocol_downloaded_at"] is not None
    detail = client.get(f"/api/checks/{ready_check}").json()
    assert detail["protocol_format"] == "docx" and detail["contract_number"] == "15/2026"

    # После нового анализа прежний выпуск не считается: протокол собирается из нового результата.
    provider.responses.append(contract_answer())
    assert client.post(f"/api/checks/{ready_check}/analyze").status_code == 202
    [item] = client.get("/api/checks").json()
    assert item["status"] == "analyzing" and item["risk_count"] is None and item["protocol_format"] is None
    run_next_job(db_factory, settings)
    [item] = client.get("/api/checks").json()
    assert item["status"] == "ready" and item["contract_number"] == "12/2026"
    assert item["protocol_downloaded_at"] is None


def test_audit_trail(client, db_factory, ready_check) -> None:
    client.delete(f"/api/checks/{ready_check}")
    entries = audit_rows(db_factory)
    assert [entry.action for entry in entries] == ["upload", "analysis_done", "delete"]
    done = entries[1]
    assert (done.model, done.tokens_in, done.tokens_out, done.cost_rub) == ("fake-model", 1000, 200, 0.5)
    assert len(done.prompt_version or "") == 12
    assert done.counterparty == "АО «Пример»" and done.contract_number == "12/2026"
    assert all(entry.check_id == ready_check and entry.user_id == 1 for entry in entries)


# --- ограничение входа -----------------------------------------------------------------------------


def test_login_attempts_are_limited(client, db_factory) -> None:
    add_user(db_factory)
    add_user(db_factory, login="second")
    for _ in range(5):
        assert (
            client.post("/api/auth/login", json={"login": "manager", "password": "wrong"}).status_code == 401
        )
    blocked = client.post("/api/auth/login", json={"login": "manager", "password": PASSWORD})
    assert blocked.status_code == 429
    # Блокировка — на 15 минут; под нагрузкой к ответу проходит секунда, и округление даёт 14.
    assert re.search("Повторите через 1[45] мин", blocked.json()["detail"])
    assert int(blocked.headers["retry-after"]) > 0
    # Другой логин с того же адреса входит: лимит по адресу выше.
    assert client.post("/api/auth/login", json={"login": "second", "password": PASSWORD}).status_code == 200


def test_login_limiter_window(monkeypatch) -> None:
    clock = [1000.0]
    monkeypatch.setattr("app.api.ratelimit.time.monotonic", lambda: clock[0])
    limiter = LoginLimiter(max_per_login=2, max_per_address=3, window_seconds=60)
    limiter.failed("a", "ip1")
    limiter.failed("a", "ip1")
    assert limiter.retry_after("a", "ip1") == 60
    clock[0] += 61
    assert limiter.retry_after("a", "ip1") is None
    limiter.failed("a", "ip1")
    limiter.succeeded("a")
    assert limiter.retry_after("a", "ip1") is None
    for name in ("b", "c", "d"):
        limiter.failed(name, "ip2")
    assert limiter.retry_after("e", "ip2") is not None


# --- команды CLI для данных организации --------------------------------------------------------------


def test_cli_organization(settings, db_factory, tmp_path, monkeypatch, capsys) -> None:
    monkeypatch.setattr(cli, "get_settings", lambda: settings)
    cli.main(
        [
            "org-set",
            "--full-name",
            "Индивидуальный предприниматель Петров Пётр Петрович",
            "--short-name",
            "ИП Петров Пётр Петрович",
            "--signer",
            "П.П. Петров",
            "--city",
            "г. Пример",
            "--requisite",
            "inn=770000000156",
        ]
    )
    image = tmp_path / "signature.png"
    image.write_bytes(png())
    cli.main(["org-asset", "signature", str(image)])
    with pytest.raises(SystemExit, match="не картинка"):
        bad = tmp_path / "bad.png"
        bad.write_bytes(b"not an image")
        cli.main(["org-asset", "stamp", str(bad)])
    with pytest.raises(SystemExit, match="Неизвестные реквизиты"):
        cli.main(["org-set", "--requisite", "passport=1"])
    # Формат реквизитов команда проверяет так же, как админка.
    with pytest.raises(SystemExit, match="ИНН: нужно 10 или 12 цифр"):
        cli.main(["org-set", "--requisite", "inn=12"])
    capsys.readouterr()
    cli.main(["org-show"])
    out = capsys.readouterr().out
    assert "Подписант: П.П. Петров" in out and "inn: 770000000156" in out
    assert "signature: " in out and "stamp: нет" in out and "Готово для протокола" in out
    cli.main(["org-asset", "signature", "--clear"])
    with db_factory() as db:
        organization = db.scalars(select(Organization)).one()
        assert organization.signature_png is None and organization.default_city == "г. Пример"


def test_saved_state_is_validated_model() -> None:
    state = draft_from_result({"answer": answer()}, BLOCKS, None, date(2026, 10, 16))
    restored = ProtocolState.model_validate(state.model_dump(mode="json"))
    assert restored == state
