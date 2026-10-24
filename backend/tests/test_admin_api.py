"""Этап 5: права на /api/admin, пользователи, версии промпта, шаблон протокола, лимиты, журнал."""

import io
import re
from datetime import timedelta

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.analysis.default_prompt import DEFAULT_ADMIN_PROMPT
from app.api.middleware import CSRF_HEADER, CSRF_HEADER_VALUE
from app.core.db import utcnow
from app.models import AuditEntry, Check, PromptVersion, UserRole, UserSession
from app.protocol.templates import NBSP
from tests.conftest import PASSWORD, add_user
from tests.test_analysis_job import add_profile, contract_answer, provider, uploaded  # noqa: F401
from tests.test_checks_api import docx_bytes, login, run_next_job, upload
from tests.test_protocol_api import add_organization

NEW_PASSWORD = "another-long-password"
PATH_VALUES = {
    "user_id": "1",
    "version_id": "1",
    "profile_id": "1",
    "run_id": "x",
    "kind": "signature",
}


@pytest.fixture
def admin(client: TestClient, db_factory) -> int:
    user_id = add_user(db_factory, login="admin", role=UserRole.ADMIN)
    login(client, "admin")
    return user_id


def admin_routes(app: FastAPI) -> list[tuple[str, str]]:
    routes = []
    for path, operations in app.openapi()["paths"].items():
        if path.startswith("/api/admin"):
            concrete = re.sub(r"\{(\w+)\}", lambda match: PATH_VALUES[match.group(1)], path)
            routes += [(method.upper(), concrete) for method in sorted(operations)]
    return routes


def audit_actions(db_factory) -> list[tuple[str, str | None]]:
    with db_factory() as db:
        return [(row.action, row.details) for row in db.scalars(select(AuditEntry).order_by(AuditEntry.id))]


# --- права --------------------------------------------------------------------------------------


def test_admin_api_requires_admin_role(app: FastAPI, client: TestClient, db_factory) -> None:
    routes = admin_routes(app)
    assert len(routes) >= 30
    for method, path in routes:
        assert client.request(method, path).status_code == 401, (method, path)
    add_user(db_factory)
    login(client, "manager")
    for method, path in routes:
        assert client.request(method, path).status_code == 403, (method, path)


# --- пользователи -------------------------------------------------------------------------------


def test_create_list_and_validate_users(client: TestClient, admin: int, db_factory) -> None:
    response = client.post(
        "/api/admin/users", json={"login": " Ivanov ", "name": "Иван Иванов", "password": PASSWORD}
    )
    assert response.status_code == 201, response.text
    created = response.json()
    assert (created["login"], created["role"], created["is_active"]) == ("ivanov", "manager", True)
    assert created["last_seen_at"] is None

    assert client.post("/api/admin/users", json={"login": "IVANOV", "password": PASSWORD}).status_code == 409
    bad_login = client.post("/api/admin/users", json={"login": "a b", "password": PASSWORD})
    assert bad_login.status_code == 400 and "Логин" in bad_login.json()["detail"]
    short = client.post("/api/admin/users", json={"login": "petrov", "password": "123"})
    assert short.status_code == 400 and "не короче" in short.json()["detail"]

    users = {user["login"]: user for user in client.get("/api/admin/users").json()}
    assert set(users) == {"admin", "ivanov"}
    assert users["admin"]["last_seen_at"] is not None  # администратор вошёл
    assert ("user_create", "ivanov, роль: менеджер") in audit_actions(db_factory)

    # Новый пользователь входит со своим паролем.
    with TestClient(client.app) as other:
        other.headers[CSRF_HEADER] = CSRF_HEADER_VALUE
        login(other, "ivanov")


def test_block_role_and_self_protection(client: TestClient, admin: int, db_factory) -> None:
    manager_id = add_user(db_factory)
    with TestClient(client.app) as manager:
        manager.headers[CSRF_HEADER] = CSRF_HEADER_VALUE
        login(manager, "manager")
        assert manager.get("/api/me").status_code == 200

        blocked = client.patch(
            f"/api/admin/users/{manager_id}", json={"is_active": False, "name": "Менеджер"}
        )
        assert blocked.status_code == 200 and blocked.json()["is_active"] is False
        # Блокировка сразу завершает сессии.
        assert manager.get("/api/me").status_code == 401
        response = manager.post("/api/auth/login", json={"login": "manager", "password": PASSWORD})
        assert response.status_code == 403

    assert client.patch(f"/api/admin/users/{manager_id}", json={"is_active": True}).json()["is_active"]
    promoted = client.patch(f"/api/admin/users/{manager_id}", json={"role": "admin"}).json()
    assert promoted["role"] == "admin"
    details = [details for action, details in audit_actions(db_factory) if action == "user_update"]
    assert details == [
        "manager: имя; заблокирован",
        "manager: разблокирован",
        "manager: роль: менеджер → администратор",
    ]

    for change in ({"is_active": False}, {"role": "manager"}):
        response = client.patch(f"/api/admin/users/{admin}", json=change)
        assert response.status_code == 409
        assert "Нельзя заблокировать себя" in response.json()["detail"]
    # Другого администратора снять можно: активный администратор остаётся.
    assert client.patch(f"/api/admin/users/{manager_id}", json={"role": "manager"}).status_code == 200
    assert client.patch("/api/admin/users/999", json={"name": "x"}).status_code == 404


def test_reset_password(client: TestClient, admin: int, db_factory) -> None:
    manager_id = add_user(db_factory)
    with TestClient(client.app) as manager:
        manager.headers[CSRF_HEADER] = CSRF_HEADER_VALUE
        login(manager, "manager")
        assert (
            client.post(f"/api/admin/users/{manager_id}/password", json={"password": "short"}).status_code
            == 400
        )
        response = client.post(f"/api/admin/users/{manager_id}/password", json={"password": NEW_PASSWORD})
        assert response.status_code == 204
        assert manager.get("/api/me").status_code == 401
        assert (
            manager.post("/api/auth/login", json={"login": "manager", "password": PASSWORD}).status_code
            == 401
        )
        assert (
            manager.post("/api/auth/login", json={"login": "manager", "password": NEW_PASSWORD}).status_code
            == 200
        )

    # Свой пароль: другие сессии администратора завершаются, текущая остаётся.
    with TestClient(client.app) as second:
        second.headers[CSRF_HEADER] = CSRF_HEADER_VALUE
        login(second, "admin")
        assert (
            client.post(f"/api/admin/users/{admin}/password", json={"password": NEW_PASSWORD}).status_code
            == 204
        )
        assert client.get("/api/me").status_code == 200
        assert second.get("/api/me").status_code == 401
    with db_factory() as db:
        assert len(db.scalars(select(UserSession).where(UserSession.user_id == admin)).all()) == 1
    assert [action for action, _ in audit_actions(db_factory)].count("user_password") == 2


# --- промпт -------------------------------------------------------------------------------------


def test_first_version_is_customer_prompt(client: TestClient, admin: int) -> None:
    versions = client.get("/api/admin/prompts").json()
    assert len(versions) == 1
    first = versions[0]
    assert first["is_active"] and first["author"] is None and first["length"] == len(DEFAULT_ADMIN_PROMPT)
    full = client.get(f"/api/admin/prompts/{first['id']}").json()
    assert full["text"] == DEFAULT_ADMIN_PROMPT
    assert client.get("/api/admin/prompts").json() == versions  # вторая версия сама не появляется
    assert client.get("/api/admin/prompts/999").status_code == 404


def test_new_version_is_used_by_analysis_and_rollback(
    client: TestClient,
    admin: int,
    db_factory,
    settings,
    tmp_path,
    provider,  # noqa: F811
) -> None:
    first_id = client.get("/api/admin/prompts").json()[0]["id"]
    response = client.post(
        "/api/admin/prompts", json={"text": "  Новый промпт: ищи риски поставщика.  ", "comment": "короче"}
    )
    assert response.status_code == 201, response.text
    second = response.json()
    assert (
        second["is_active"]
        and second["author"] == "admin"
        and second["text"] == "Новый промпт: ищи риски поставщика."
    )
    same = client.post("/api/admin/prompts", json={"text": "Новый промпт: ищи риски поставщика."})
    assert same.status_code == 400 and "не отличается" in same.json()["detail"]
    draft = client.post("/api/admin/prompts", json={"text": "Черновик", "activate": False}).json()
    assert draft["is_active"] is False

    versions = client.get("/api/admin/prompts").json()
    assert [version["id"] for version in versions] == [draft["id"], second["id"], first_id]
    assert [version["is_active"] for version in versions] == [False, True, False]

    add_profile(db_factory)
    provider.responses.extend([contract_answer(), contract_answer()])
    check_id = uploaded(client, db_factory, settings, tmp_path)
    run_next_job(db_factory, settings)
    analysis = client.get(f"/api/checks/{check_id}").json()["analysis"]
    assert analysis["prompt_version_id"] == second["id"]
    assert analysis["prompt_fingerprint"] == second["fingerprint"]
    assert provider.requests[0].system.startswith("Новый промпт: ищи риски поставщика.")

    # Откат на первую версию: следующий анализ идёт с ней.
    assert client.post(f"/api/admin/prompts/{first_id}/activate").json()["is_active"] is True
    with db_factory() as db:
        assert db.scalars(select(PromptVersion.id).where(PromptVersion.is_active.is_(True))).all() == [
            first_id
        ]
    assert client.post(f"/api/checks/{check_id}/analyze").status_code == 202
    run_next_job(db_factory, settings)
    assert provider.requests[1].system.startswith(DEFAULT_ADMIN_PROMPT.strip()[:200])
    actions = audit_actions(db_factory)
    assert ("prompt_activate", f"версия {second['id']} → {first_id}") in actions
    with db_factory() as db:
        done = db.scalars(select(AuditEntry).where(AuditEntry.action == "analysis_done")).all()
        assert [row.prompt_version for row in done] == [second["fingerprint"], versions[2]["fingerprint"]]


# --- шаблон протокола ---------------------------------------------------------------------------


def test_protocol_template_validation_and_use(
    client: TestClient,
    admin: int,
    db_factory,
    settings,
    tmp_path,
    provider,  # noqa: F811
) -> None:
    body = client.get("/api/admin/protocol-template").json()
    assert body["template"] == body["defaults"]
    assert "{first_party}" in body["template"]["preamble"]
    assert set(body["fields"]) == {"preamble", "closing", "appendices"}

    template = {**body["template"], "preamble": "{first_party} и {secret.__class__}"}
    response = client.put("/api/admin/protocol-template", json=template)
    assert response.status_code == 422
    assert "Неизвестная подстановка" in str(response.json()["detail"])
    for preamble in ("{0} и {first_party}", "{first_party!r}", "{first_party:>10}", "{first_party"):
        assert (
            client.put("/api/admin/protocol-template", json={**template, "preamble": preamble}).status_code
            == 422
        )

    template = {
        **body["template"],
        "title": "ПРОТОКОЛ СОГЛАСОВАНИЯ РАЗНОГЛАСИЙ",
        "preamble": "{first_party} и {second_party} составили протокол к {contract}:",
        "closing_items": ["Прочие условия Договора № {contract_number} остаются в силе.", "  "],
        "include_agreed_column": False,
        "table_font_size": 9,
    }
    saved = client.put("/api/admin/protocol-template", json=template)
    assert saved.status_code == 200, saved.text
    assert saved.json()["closing_items"] == ["Прочие условия Договора № {contract_number} остаются в силе."]
    assert ("template_update", None) in audit_actions(db_factory)

    preview = client.post("/api/admin/protocol-template/preview", json=template).json()
    assert preview["title"] == "ПРОТОКОЛ СОГЛАСОВАНИЯ РАЗНОГЛАСИЙ"
    assert preview["preamble"].startswith("Общество с ограниченной ответственностью «Образец»")
    assert preview["closing_items"][0] == "Прочие условия Договора № 1/2026 остаются в силе."

    # Новая проверка получает шаблон в черновик протокола.
    add_organization(db_factory)
    add_profile(db_factory)
    provider.responses.append(contract_answer())
    check_id = uploaded(client, db_factory, settings, tmp_path)
    run_next_job(db_factory, settings)
    protocol = client.get(f"/api/checks/{check_id}/protocol").json()
    meta = protocol["state"]["meta"]
    assert (meta["title"], meta["include_agreed_column"], meta["table_font_size"]) == (
        "ПРОТОКОЛ СОГЛАСОВАНИЯ РАЗНОГЛАСИЙ",
        False,
        9,
    )
    assert protocol["preamble_preview"].endswith(
        f"составили протокол к Договору поставки №{NBSP}12/2026 от 31.07.2026:"
    )


def test_protocol_template_preview_docx(client: TestClient, admin: int, db_factory) -> None:
    add_organization(db_factory)
    template = client.get("/api/admin/protocol-template").json()["template"]
    response = client.post(
        "/api/admin/protocol-template/preview.docx?signature=true&stamp=false", json=template
    )
    assert response.status_code == 200
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    document = Document(io.BytesIO(response.content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "ПРОТОКОЛ РАЗНОГЛАСИЙ" in text
    assert "Индивидуальный предприниматель Петров Пётр Петрович" in text
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    assert any("0,1 %" in cell for cell in cells)


# --- лимиты ---------------------------------------------------------------------------------------


def test_limits_apply_to_new_uploads(client: TestClient, admin: int, db_factory, settings, tmp_path) -> None:
    body = client.get("/api/admin/settings").json()
    assert body == {
        "check_ttl_hours": 24,
        "max_upload_mb": 50,
        "defaults": {"check_ttl_hours": 24, "max_upload_mb": 50},
    }
    assert (
        client.put("/api/admin/settings", json={"check_ttl_hours": 0, "max_upload_mb": 1}).status_code == 422
    )
    saved = client.put("/api/admin/settings", json={"check_ttl_hours": 2, "max_upload_mb": 1})
    assert saved.status_code == 200 and saved.json()["check_ttl_hours"] == 2
    assert client.get("/api/limits").json() == {"check_ttl_hours": 2, "max_upload_mb": 1}

    response = upload(client, docx_bytes(tmp_path, "1.1. Текст"))
    assert response.status_code == 201
    created = response.json()
    with db_factory() as db:
        check = db.get(Check, created["id"])
        assert check is not None
        assert timedelta(hours=1, minutes=59) < check.expires_at - utcnow() <= timedelta(hours=2)

    big = b"PK" + b"0" * (1024 * 1024 + 10)
    response = client.post("/api/checks", files={"file": ("big.docx", big, "application/octet-stream")})
    assert response.status_code == 413 and "1 МБ" in response.json()["detail"]
    # Запрос больше лимита с запасом отклоняется ещё до чтения тела.
    huge = b"0" * (2 * 1024 * 1024 + 10)
    response = client.post("/api/checks", files={"file": ("huge.docx", huge, "application/octet-stream")})
    assert response.status_code == 413
    assert any(action == "settings_update" for action, _ in audit_actions(db_factory))


# --- журнал ---------------------------------------------------------------------------------------


def test_audit_filters_totals_and_csv(client: TestClient, admin: int, db_factory) -> None:
    manager_id = add_user(db_factory)
    now = utcnow()
    with db_factory() as db:
        db.add_all(
            [
                AuditEntry(at=now - timedelta(days=2), user_id=manager_id, action="upload", check_id="c1"),
                AuditEntry(
                    at=now - timedelta(hours=1),
                    user_id=manager_id,
                    action="analysis_done",
                    check_id="c1",
                    counterparty='=HYPERLINK("x")',
                    contract_number="12/2026",
                    model="vendor/model",
                    tokens_in=1000,
                    tokens_out=200,
                    cost_rub=1.25,
                ),
                AuditEntry(
                    at=now,
                    user_id=manager_id,
                    action="protocol_docx",
                    check_id="c1",
                    with_signature=True,
                    with_stamp=False,
                ),
            ]
        )
        db.commit()
    page = client.get("/api/admin/audit", params={"user_id": manager_id}).json()
    assert page["total"] == 3
    assert [item["action"] for item in page["items"]] == ["protocol_docx", "analysis_done", "upload"]
    assert page["items"][0]["login"] == "manager" and page["items"][0]["action_label"] == "протокол Word"
    assert page["totals"] == {"count": 3, "tokens_in": 1000, "tokens_out": 200, "cost_rub": 1.25}
    assert page["actions"]["prompt_test"] == "прогон черновика промпта"

    since = (now - timedelta(days=1)).isoformat()
    recent = client.get("/api/admin/audit", params={"user_id": manager_id, "at_from": since}).json()
    assert recent["total"] == 2
    only = client.get("/api/admin/audit", params=[("action", "upload"), ("action", "protocol_docx")]).json()
    assert {item["action"] for item in only["items"]} == {"upload", "protocol_docx"}
    paged = client.get("/api/admin/audit", params={"user_id": manager_id, "limit": 1, "offset": 1}).json()
    assert [item["action"] for item in paged["items"]] == ["analysis_done"] and paged["total"] == 3

    response = client.get("/api/admin/audit.csv", params={"user_id": manager_id, "utc_offset": 180})
    assert response.status_code == 200
    assert response.content.startswith(b"\xef\xbb\xbf")
    lines = response.content.decode("utf-8-sig").splitlines()
    assert lines[0].startswith("Дата и время;Пользователь;Действие")
    assert len(lines) == 4
    done = next(line for line in lines if "анализ выполнен" in line)
    # Текст из договора не превращается в формулу Excel.
    assert "'=HYPERLINK" in done and ";=HYPERLINK" not in done and ";1,2500;" in done
    expected_time = (now - timedelta(hours=1) + timedelta(hours=3)).strftime("%Y-%m-%d %H:%M")
    assert done.startswith(expected_time)
