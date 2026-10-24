"""Сборка docx протокола разногласий по documentation/PROTOCOL_FORMAT.md.

python-docx создаёт документ, а то, чего в нём нет (повтор шапки таблицы, заливка, плавающие картинки,
поле номера страницы), дописывается прямо в XML.
"""

import io
from copy import deepcopy

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.section import _Footer, _Header
from docx.shared import Emu, Length, Mm, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from app.protocol import assets
from app.protocol.model import ProtocolData, ProtocolImages, ProtocolRow
from app.protocol.templates import long_date

PAGE_WIDTH, PAGE_HEIGHT = Mm(297), Mm(210)
MARGIN_LEFT, MARGIN_RIGHT, MARGIN_TOP, MARGIN_BOTTOM = Mm(20), Mm(15), Mm(15), Mm(15)
HEADER_FOOTER_DISTANCE = Mm(8)
TEXT_WIDTH = Emu(PAGE_WIDTH - MARGIN_LEFT - MARGIN_RIGHT)
MAX_HEADER_IMAGE_WIDTH = Mm(187)

FONT = "Times New Roman"
HEADER_FILL = "D9E2F3"
FOOTER_COLOR = RGBColor(0x59, 0x59, 0x59)
# Доли ширины текста, как в образце; «Пункт Договора» чуть шире, чтобы слово «Договора» не переносилось.
COLUMNS_WITH_AGREED = (3.5, 7.5, 23.5, 36.0, 29.5)
COLUMNS_WITHOUT_AGREED = (4.0, 9.0, 37.0, 50.0)
HEADERS = ("№", "Пункт Договора", "Редакция Покупателя", "Редакция Поставщика", "Согласованная редакция")
# В «Согласованной редакции» эти слова выделяются, как в образце.
SHORT_AGREED = {"Редакция Поставщика", "Редакция Покупателя"}
SIGNATURE_COLUMNS = (65.0, 35.0)
SIGNATURE_LINE = "________________"
SIGNATURE_WIDTH = Mm(34)
SIGNATURE_LINE_HEIGHT = Pt(12)
SIGNATURE_SPACE_LINES = 3
STAMP_WIDTH = Mm(40)
CELL_MARGIN_TWIPS = 85  # 1,5 мм


def build_docx(
    data: ProtocolData,
    images: ProtocolImages | None = None,
    *,
    with_signature: bool = False,
    with_stamp: bool = False,
) -> bytes:
    images = images or ProtocolImages()
    document = Document()
    _setup_styles(document)
    _setup_section(document, data, images)
    _title_block(document, data)
    _rows_table(document, data)
    _closing(document, data)
    _signatures(
        document,
        data,
        signature=images.signature if with_signature else None,
        stamp=images.stamp if with_stamp else None,
    )
    properties = document.core_properties
    properties.title = data.title
    properties.author = data.supplier.name
    properties.last_modified_by = data.supplier.name
    properties.comments = ""
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- оформление -------------------------------------------------------------------------


def _set_fonts(rpr) -> None:
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    # Темы шрифтов (asciiTheme и др.) важнее явных имён: убираем их, иначе Word возьмёт шрифт темы.
    for attribute in list(fonts.attrib):
        if attribute.endswith("Theme"):
            del fonts.attrib[attribute]
    for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attribute), FONT)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "ru-RU")
    lang.set(qn("w:eastAsia"), "ru-RU")


def _setup_styles(document: DocxDocument) -> None:
    defaults = document.styles.element.find(qn("w:docDefaults"))
    if defaults is not None:
        rpr = defaults.find(f"{qn('w:rPrDefault')}/{qn('w:rPr')}")
        if rpr is not None:
            _set_fonts(rpr)
    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    _set_fonts(normal.element.get_or_add_rPr())
    paragraph_format = normal.paragraph_format
    paragraph_format.space_before = Pt(0)
    paragraph_format.space_after = Pt(0)
    paragraph_format.line_spacing = 1.0


def _setup_section(document: DocxDocument, data: ProtocolData, images: ProtocolImages) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = PAGE_WIDTH, PAGE_HEIGHT
    section.left_margin, section.right_margin = MARGIN_LEFT, MARGIN_RIGHT
    section.top_margin, section.bottom_margin = MARGIN_TOP, MARGIN_BOTTOM
    section.header_distance = section.footer_distance = HEADER_FOOTER_DISTANCE

    letterhead = data.letterhead.enabled
    if letterhead and images.header:
        # Шапка бланка — только на первой странице.
        section.different_first_page_header_footer = True
        _header_image(section.first_page_header, images.header)
        _footer(section.first_page_footer, data.letterhead.footer)
        _footer(section.footer, data.letterhead.footer)
    else:
        _footer(section.footer, data.letterhead.footer if letterhead else "")


def _header_image(header: _Header, image: bytes) -> None:
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    width = min(TEXT_WIDTH, MAX_HEADER_IMAGE_WIDTH)
    paragraph.add_run().add_picture(io.BytesIO(image), width=width)


def _footer(footer: _Footer, contacts: str) -> None:
    first = footer.paragraphs[0]
    if contacts.strip():
        first.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(first, contacts.strip(), 8).font.color.rgb = FOOTER_COLOR
        page = footer.add_paragraph()
    else:
        page = first
    page.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(page, "Страница ", 9)
    _page_field(page, 9)


def _page_field(paragraph: Paragraph, size: float) -> None:
    """Номер страницы — сложное поле PAGE, размер задан у каждой его части.

    Word, обновляя поле, оформляет номер по первому символу кода поля. У простого поля (w:fldSimple)
    кода как отдельного прогона нет, и Word набирал номер размером стиля Normal (11 pt вместо 9).
    """
    for part in ("begin", " PAGE ", "separate", "1", "end"):
        run = _run(paragraph, "", size)
        if part in ("begin", "separate", "end"):
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), part)
        else:
            node = OxmlElement("w:instrText" if part == " PAGE " else "w:t")
            node.set(qn("xml:space"), "preserve")
            node.text = part
        run._r.append(node)


def _run(paragraph: Paragraph, text: str, size: float, bold: bool = False):
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    return run


def _paragraph(
    container,
    text: str,
    size: float,
    *,
    bold: bool = False,
    align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT,
    before: float = 0,
    after: float = 0,
    first_line: Length | None = None,
    line_spacing: float = 1.0,
    keep_with_next: bool = False,
) -> Paragraph:
    paragraph = container.add_paragraph()
    paragraph.alignment = align
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_before = Pt(before)
    paragraph_format.space_after = Pt(after)
    paragraph_format.line_spacing = line_spacing
    paragraph_format.keep_with_next = keep_with_next
    if first_line is not None:
        paragraph_format.first_line_indent = first_line
    if text:
        _run(paragraph, text, size, bold)
    return paragraph


# --- части документа ----------------------------------------------------------------------


def _title_block(document: DocxDocument, data: ProtocolData) -> None:
    # Документ python-docx начинается с пустого абзаца: используем его под заголовок.
    title = document.paragraphs[0] if document.paragraphs else document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    _run(title, data.title, 16, bold=True)
    if data.subtitle:
        _paragraph(document, data.subtitle, 14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)

    place = _paragraph(document, "", 11, after=12)
    place.paragraph_format.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    _run(place, data.city, 11)
    if data.date:
        _run(place, "\t" + long_date(data.date), 11)

    if data.preamble:
        _paragraph(
            document,
            data.preamble,
            11,
            align=WD_ALIGN_PARAGRAPH.JUSTIFY,
            after=10,
            first_line=Mm(12.5),
            line_spacing=1.15,
        )


def _rows_table(document: DocxDocument, data: ProtocolData) -> None:
    columns = COLUMNS_WITH_AGREED if data.include_agreed_column else COLUMNS_WITHOUT_AGREED
    headers = HEADERS if data.include_agreed_column else HEADERS[:4]
    widths = [Emu(int(TEXT_WIDTH * share / 100)) for share in columns]

    table = document.add_table(rows=1, cols=len(columns))
    table.style = document.styles["Table Grid"]
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _fixed_layout(table, widths)

    header = table.rows[0]
    _row_properties(header._tr, repeat_header=True, cant_split=True)
    size = data.table_font_size
    for cell, text, width in zip(header.cells, headers, widths, strict=True):
        cell.width = width
        _shade(cell, HEADER_FILL)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(paragraph, text, size, bold=True)

    for number, row in enumerate(data.rows, start=1):
        cells = table.add_row().cells
        for cell, width in zip(cells, widths, strict=True):
            cell.width = width
        _centered(cells[0], str(number), size)
        _centered(cells[1], row.clause, size)
        _text_cell(cells[2], row.buyer_text, size)
        _text_cell(cells[3], row.supplier_text, size)
        if data.include_agreed_column:
            _agreed_cell(cells[4], row, size)


def _centered(cell: _Cell, text: str, size: float) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(paragraph, text, size, bold=True)


def _text_cell(cell: _Cell, text: str, size: float) -> None:
    parts = [part.strip() for part in text.replace("\r\n", "\n").split("\n") if part.strip()]
    first = True
    for part in parts:
        paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.space_after = Pt(6)
        _run(paragraph, part, size)
        first = False
    if parts:
        # У последнего абзаца ячейки отступ после не нужен.
        cell.paragraphs[-1].paragraph_format.space_after = Pt(0)


def _agreed_cell(cell: _Cell, row: ProtocolRow, size: float) -> None:
    text = row.agreed_text.strip()
    if text in SHORT_AGREED:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _centered(cell, text, size)
    else:
        _text_cell(cell, text, size)


def _closing(document: DocxDocument, data: ProtocolData) -> None:
    if not data.closing_items:
        return
    _paragraph(document, data.closing_title, 13, bold=True, before=14, after=6, keep_with_next=True)
    last = len(data.closing_items)
    for number, item in enumerate(data.closing_items, start=1):
        # Последний пункт держится вместе с подписями: подписи не остаются одни на странице.
        _paragraph(
            document,
            f"{number}. {item.strip()}",
            11,
            align=WD_ALIGN_PARAGRAPH.JUSTIFY,
            after=4,
            keep_with_next=number == last,
        )


def _signatures(
    document: DocxDocument, data: ProtocolData, signature: bytes | None, stamp: bytes | None
) -> None:
    _paragraph(document, "", 11, after=18, keep_with_next=True)
    widths = [Emu(int(TEXT_WIDTH * share / 100)) for share in SIGNATURE_COLUMNS]
    table = document.add_table(rows=1, cols=2)
    _fixed_layout(table, widths)
    _row_properties(table.rows[0]._tr, repeat_header=False, cant_split=True)
    buyer_cell, supplier_cell = table.rows[0].cells
    for cell, width in zip((buyer_cell, supplier_cell), widths, strict=True):
        cell.width = width

    _party(buyer_cell, data.buyer.heading, data.buyer.name, data.buyer.signer)
    line, seal = _party(supplier_cell, data.supplier.heading, data.supplier.name, data.supplier.signer)
    # Смещения считаются от верха абзаца. У абзацев блока подписей нет интервалов, а высота строк точная
    # (SIGNATURE_LINE_HEIGHT): так положение картинок одинаково в Word и LibreOffice.
    if signature:
        picture, ratio = assets.prepare_signature(signature)
        height = Emu(int(SIGNATURE_WIDTH * ratio))
        # Подпись стоит на линии: низ росчерка — чуть ниже строки, верх — под наименованием поставщика.
        _floating_picture(line, picture, SIGNATURE_WIDTH, x=Mm(2), y=Emu(Mm(8.5) - height), z=251658240)
    if stamp:
        picture, ratio = assets.prepare_stamp(stamp, data.contract_number or data.title)
        height = Emu(int(STAMP_WIDTH * ratio))
        # Печать — на «М.П.»: верх оттиска примерно на линии подписи, он заходит только на нижний край
        # росчерка, поэтому подпись видна, а наименование поставщика не закрыто.
        _floating_picture(seal, picture, STAMP_WIDTH, x=Mm(0), y=Emu(Mm(12) - height // 2), z=251659264)


def _party(cell: _Cell, heading: str, name: str, signer: str) -> tuple[Paragraph, Paragraph]:
    first = cell.paragraphs[0]
    _run(first, heading, 11, bold=True)
    if name:
        _run(cell.add_paragraph(), name, 11)
    # Место под подпись — пустые строки той же точной высоты: геометрия блока не зависит от программы.
    for _ in range(SIGNATURE_SPACE_LINES):
        cell.add_paragraph()
    line = cell.add_paragraph()
    _run(line, f"{SIGNATURE_LINE} / {signer or SIGNATURE_LINE} /", 11)
    cell.add_paragraph()
    seal = cell.add_paragraph()
    _run(seal, "М.П.", 11)
    for paragraph in cell.paragraphs:
        paragraph_format = paragraph.paragraph_format
        paragraph_format.keep_with_next = True
        paragraph_format.space_before = Pt(0)
        paragraph_format.space_after = Pt(0)
        paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        paragraph_format.line_spacing = SIGNATURE_LINE_HEIGHT
    return line, seal


# --- XML ----------------------------------------------------------------------------------


def _fixed_layout(table: Table, widths: list[Emu]) -> None:
    """Фиксированные ширины колонок: Word и LibreOffice не пересчитывают их под содержимое.

    Порядок элементов в w:tblPr задан схемой; Word считает документ повреждённым, если его нарушить.
    """
    table.autofit = False  # python-docx сам ставит w:tblLayout type="fixed" на нужное место
    tbl_pr = table._tbl.tblPr
    tbl_width = tbl_pr.find(qn("w:tblW"))
    if tbl_width is None:
        tbl_width = OxmlElement("w:tblW")
        tbl_pr.append(tbl_width)
    tbl_width.set(qn("w:w"), str(int(sum(widths) / 635)))  # EMU → twips
    tbl_width.set(qn("w:type"), "dxa")
    margins = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        margin = OxmlElement(f"w:{side}")
        margin.set(qn("w:w"), str(CELL_MARGIN_TWIPS))
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    # w:tblCellMar идёт после w:tblLayout и перед w:tblLook.
    following = next(
        (
            tbl_pr.find(qn(tag))
            for tag in ("w:tblLook", "w:tblCaption", "w:tblDescription")
            if tbl_pr.find(qn(tag)) is not None
        ),
        None,
    )
    if following is not None:
        following.addprevious(margins)
    else:
        tbl_pr.append(margins)
    for column, width in zip(table.columns, widths, strict=True):
        column.width = width


def _row_properties(tr, *, repeat_header: bool, cant_split: bool) -> None:
    tr_pr = tr.get_or_add_trPr()
    if cant_split:
        tr_pr.append(OxmlElement("w:cantSplit"))
    if repeat_header:
        tr_pr.append(OxmlElement("w:tblHeader"))


def _shade(cell: _Cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _floating_picture(
    paragraph: Paragraph, picture: bytes, width: Length, *, x: Length, y: Length, z: int
) -> None:
    """Картинка поверх текста без обтекания, привязанная к абзацу (wp:anchor вместо wp:inline)."""
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    run.add_picture(io.BytesIO(picture), width=width)
    inline = run._r.findall(f".//{qn('wp:inline')}")[-1]

    anchor = OxmlElement("wp:anchor")
    for name, value in {
        "distT": "0", "distB": "0", "distL": "0", "distR": "0", "simplePos": "0",
        "relativeHeight": str(z), "behindDoc": "0", "locked": "0", "layoutInCell": "1", "allowOverlap": "1",
    }.items():  # fmt: skip
        anchor.set(name, value)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    anchor.append(simple)
    for tag, relative, offset in (("wp:positionH", "column", x), ("wp:positionV", "paragraph", y)):
        position = OxmlElement(tag)
        position.set("relativeFrom", relative)
        value = OxmlElement("wp:posOffset")
        value.text = str(int(offset))
        position.append(value)
        anchor.append(position)
    anchor.append(deepcopy(inline.find(qn("wp:extent"))))
    effect = OxmlElement("wp:effectExtent")
    for side in ("l", "t", "r", "b"):
        effect.set(side, "0")
    anchor.append(effect)
    anchor.append(OxmlElement("wp:wrapNone"))
    for tag in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
        anchor.append(deepcopy(inline.find(qn(tag))))
    inline.getparent().replace(inline, anchor)
